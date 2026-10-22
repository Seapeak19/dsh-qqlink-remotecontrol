import sys
import os
import json
import shutil

# 项目本地 PDF 库（不动 DSH 运行时）
_LIB = os.path.join(os.path.dirname(__file__), "pylibs")
if os.path.isdir(_LIB) and _LIB not in sys.path:
    sys.path.append(_LIB)

sys.stdout.reconfigure(encoding="utf-8")


def out(obj):
    print(json.dumps(obj, ensure_ascii=False))
    sys.exit(0)


def fail(msg):
    out({"ok": False, "error": str(msg)})


def load_args():
    with open(sys.argv[2], "r", encoding="utf-8") as f:
        return json.load(f)


def parse_pages(spec, total):
    if not spec:
        return list(range(total))
    idx = []
    for part in str(spec).replace("，", ",").split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a1, b1 = part.split("-", 1)
            try:
                start, end = int(a1), int(b1)
            except Exception:
                continue
            idx.extend(range(start - 1, end))
        else:
            try:
                idx.append(int(part) - 1)
            except Exception:
                pass
    return [i for i in idx if 0 <= i < total]


def write_pdf(writer, src, save_as):
    if os.path.abspath(src).lower() == os.path.abspath(save_as).lower():
        shutil.copy2(src, save_as + ".bak")
    tmp = save_as + ".tmp"
    writer.write(tmp)
    os.replace(tmp, save_as)


action = sys.argv[1]

try:
    a = load_args()

    if action == "docx_read":
        from docx import Document
        p = a["path"]
        if not os.path.exists(p):
            fail("文件不存在")
        doc = Document(p)
        paras = [x.text for x in doc.paragraphs]
        tables = []
        for t in doc.tables:
            tables.append([[c.text for c in row.cells] for row in t.rows])
        text = "\n".join([x for x in paras if x.strip()])
        out({
            "ok": True, "path": p,
            "paragraphs": paras, "tables": tables,
            "text": text[:20000], "paragraphCount": len(paras),
        })

    if action == "docx_edit":
        from docx import Document
        p = a["path"]
        find = a.get("find") or ""
        replace = a.get("replace") or ""
        append = a.get("append") or ""
        save_as = a.get("saveAs") or p
        if not os.path.exists(p):
            fail("文件不存在")
        if os.path.abspath(p).lower() == os.path.abspath(save_as).lower():
            shutil.copy2(p, p + ".bak")
        doc = Document(p)
        changed = 0
        if find:
            for para in doc.paragraphs:
                if find in para.text:
                    para.text = para.text.replace(find, replace)
                    changed += 1
            for t in doc.tables:
                for row in t.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            if find in para.text:
                                para.text = para.text.replace(find, replace)
                                changed += 1
        if append:
            for line in append.splitlines():
                doc.add_paragraph(line)
        doc.save(save_as)
        out({"ok": True, "path": save_as, "changed": changed, "appended": bool(append)})

    if action == "xlsx_read":
        import openpyxl
        p = a["path"]
        if not os.path.exists(p):
            fail("文件不存在")
        wb = openpyxl.load_workbook(p, data_only=False)
        sheet = a.get("sheet")
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        maxr = min(ws.max_row or 1, 200)
        maxc = min(ws.max_column or 1, 50)
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=maxr, max_col=maxc, values_only=True):
            rows.append(["" if v is None else str(v) for v in row])
        out({
            "ok": True, "path": p,
            "sheets": wb.sheetnames, "sheet": ws.title,
            "dimensions": ws.dimensions, "rows": rows,
        })

    if action == "xlsx_edit":
        import openpyxl
        p = a["path"]
        save_as = a.get("saveAs") or p
        if not os.path.exists(p):
            fail("文件不存在")
        if os.path.abspath(p).lower() == os.path.abspath(save_as).lower():
            shutil.copy2(p, p + ".bak")
        wb = openpyxl.load_workbook(p)
        sheet = a.get("sheet")
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        changed = 0
        for cell in a.get("cells") or []:
            ref = str(cell.get("cell") or "").strip()
            if not ref:
                continue
            ws[ref] = cell.get("value")
            changed += 1
        for row in a.get("appendRows") or []:
            ws.append(row)
            changed += 1
        wb.save(save_as)
        out({"ok": True, "path": save_as, "sheet": ws.title, "changed": changed})

    if action == "pdf_info":
        from pypdf import PdfReader
        p = a["path"]
        if not os.path.exists(p):
            fail("文件不存在")
        r = PdfReader(p)
        texts = []
        for i, page in enumerate(r.pages[:50]):
            try:
                t = page.extract_text() or ""
            except Exception:
                t = ""
            texts.append({"page": i + 1, "text": t[:2000]})
        meta = {}
        try:
            for k, v in (r.metadata or {}).items():
                meta[str(k)] = str(v)
        except Exception:
            pass
        out({"ok": True, "path": p, "pages": len(r.pages), "metadata": meta, "texts": texts})

    if action == "pdf_extract":
        from pypdf import PdfReader, PdfWriter
        p = a["path"]
        save_as = a["saveAs"]
        r = PdfReader(p)
        idx = parse_pages(a.get("pages"), len(r.pages))
        if not idx:
            fail("没有匹配的页码")
        w = PdfWriter()
        for i in idx:
            w.add_page(r.pages[i])
        write_pdf(w, p, save_as)
        out({"ok": True, "path": save_as, "pages": len(idx)})

    if action == "pdf_merge":
        from pypdf import PdfReader, PdfWriter
        paths = a.get("paths") or []
        save_as = a["saveAs"]
        if len(paths) < 2:
            fail("至少需要两个 PDF")
        w = PdfWriter()
        total = 0
        for p in paths:
            r = PdfReader(p)
            for page in r.pages:
                w.add_page(page)
                total += 1
        write_pdf(w, paths[0], save_as)
        out({"ok": True, "path": save_as, "pages": total})

    if action == "pdf_rotate":
        from pypdf import PdfReader, PdfWriter
        p = a["path"]
        save_as = a["saveAs"]
        angle = int(a.get("angle") or 90)
        r = PdfReader(p)
        idx = set(parse_pages(a.get("pages"), len(r.pages)))
        w = PdfWriter()
        for i, page in enumerate(r.pages):
            w.add_page(page.rotate(angle) if i in idx else page)
        write_pdf(w, p, save_as)
        out({"ok": True, "path": save_as, "rotated": len(idx), "angle": angle})

    if action == "pdf_add_text":
        from io import BytesIO
        from pypdf import PdfReader, PdfWriter
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont

        p = a["path"]
        save_as = a["saveAs"]
        text = str(a.get("text") or "")
        page_no = int(a.get("page") or 1)
        x = float(a.get("x") or 72)
        y = float(a.get("y") or 72)
        font_size = float(a.get("fontSize") or 14)

        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        r = PdfReader(p)
        if page_no < 1 or page_no > len(r.pages):
            fail("页码超出范围")
        w = PdfWriter()
        for i, page in enumerate(r.pages):
            if i == page_no - 1:
                width = float(page.mediabox.width)
                height = float(page.mediabox.height)
                buf = BytesIO()
                c = canvas.Canvas(buf, pagesize=(width, height))
                c.setFont("STSong-Light", font_size)
                ty = height - y
                for line in text.splitlines() or [""]:
                    c.drawString(x, ty, line)
                    ty -= font_size * 1.4
                c.save()
                buf.seek(0)
                overlay = PdfReader(buf).pages[0]
                page.merge_page(overlay)
            w.add_page(page)
        write_pdf(w, p, save_as)
        out({"ok": True, "path": save_as, "page": page_no})

    if action == "pdf_text_to_docx":
        from pypdf import PdfReader
        from docx import Document
        p = a["path"]
        save_as = a["saveAs"]
        r = PdfReader(p)
        doc = Document()
        doc.add_heading("PDF 文字提取", 0)
        doc.add_paragraph("源文件：" + p)
        doc.add_paragraph("说明：这里只提取文字，不保留原始排版、图片和表格样式。")
        for i, page in enumerate(r.pages):
            doc.add_heading(f"第 {i + 1} 页", level=1)
            text = page.extract_text() or ""
            for line in text.splitlines():
                doc.add_paragraph(line)
        doc.save(save_as)
        out({"ok": True, "path": save_as, "pages": len(r.pages)})

    fail("未知动作: " + action)
except Exception as e:
    fail(e)